"""Builds tests/fixtures/word/preview-contract.docx, the sample contract of the Word preview (T-14).

Same patterns as prototypes/phase0/spike6-realistic-docx.py: python-docx, a bordered table with a merged cell, a header
and a footer with tags, a page break so the preview shows two pages, a hyperlink (the preview must not follow it).

    python3 tests/fixtures/word/make-preview-sample.py tests/fixtures/word/preview-contract.docx
"""
import sys

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

d = Document()
d.add_heading("HỢP ĐỒNG LAO ĐỘNG", 0)
p = d.add_paragraph()
p.add_run("Số: ").bold = True
p.add_run("{{so_hop_dong}}")
p = d.add_paragraph()
p.add_run("Họ và tên: ").bold = True
p.add_run("{{ho_ten}}")
d.add_paragraph("Ngày sinh: {{ngay_sinh}}; CCCD: {{cccd}}; Địa chỉ: {{dia_chi}}.")
d.add_paragraph("Ghi chú: {{ghi_chu}}")
t = d.add_table(rows=3, cols=2)
t.style = "Table Grid"
t.cell(0, 0).text = "Chức vụ"
t.cell(0, 1).text = "{{chuc_vu}}"
t.cell(1, 0).text = "Mức lương"
t.cell(1, 1).text = "{{luong}} đồng/tháng"
t.cell(2, 0).merge(t.cell(2, 1)).text = "Phòng ban: {{phong_ban}}"

# A hyperlink, written the way Word writes it.
p = d.add_paragraph("Điều khoản chi tiết: ")
rid = d.part.relate_to("https://example.com/dieu-khoan", RT.HYPERLINK, is_external=True)
link = OxmlElement("w:hyperlink")
link.set(qn("r:id"), rid)
run = OxmlElement("w:r")
text = OxmlElement("w:t")
text.text = "xem tại đây"
run.append(text)
link.append(run)
p._p.append(link)

d.add_page_break()
d.add_heading("Điều 2", level=1)
d.add_paragraph("Người lao động {{ho_ten}} làm việc tại {{phong_ban}} theo hợp đồng số {{so_hop_dong}}.")

s = d.sections[0]
s.header.paragraphs[0].text = "Công ty Sheetstamp - Hợp đồng {{so_hop_dong}}"
s.footer.paragraphs[0].text = "{{ho_ten}} - trang"
d.save(sys.argv[1])
